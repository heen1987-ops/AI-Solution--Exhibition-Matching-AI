from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from app.api.v1.routers import imports as imports_router
from app.db.session import get_db
from app.main import app
from app.schemas.imports import ExcelMatchExportRequest
from app.services import batch_matching_export as export_service
from app.services.batch_matching_export import (
    EXCEL_BATCH_MATCH_SCHEMA_VERSION,
    XLSX_MEDIA_TYPE,
    BatchMatchEntry,
    BatchMatchingReport,
    BatchProfileStatus,
    SourceProfileRef,
    render_batch_matching_workbook,
    run_batch_matching,
)
from app.services.matching.errors import personalization_disabled
from app.services.matching.types import (
    MatchCandidate,
    MatchReasonDraft,
    PipelineTrace,
    RecommendationOutcome,
)
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError

TENANT_ID = uuid.uuid4()
EVENT_ID = uuid.uuid4()
PROFILE_ID = uuid.uuid4()
USER_ID = uuid.uuid4()
EXHIBITOR_ID = uuid.uuid4()
SESSION_ID = uuid.uuid4()


def _report() -> BatchMatchingReport:
    return BatchMatchingReport(
        schema_version=EXCEL_BATCH_MATCH_SCHEMA_VERSION,
        tenant_id=TENANT_ID,
        event_id=EVENT_ID,
        source_system_code="EXCEL_UPLOAD",
        top_n=5,
        generated_at=datetime(2026, 8, 2, 12, 0, tzinfo=UTC),
        matches=(
            BatchMatchEntry(
                source_record_id="VISITOR-001",
                rank=1,
                exhibitor_name="샘플 양조장",
                exhibitor_id=str(EXHIBITOR_ID),
                object_id=str(EXHIBITOR_ID),
                score_100=87.25,
                match_level="VERY_HIGH",
                recommended_action="VISIT_NOW",
                reason_codes=("CATEGORY_MATCH", "GOAL_MATCH"),
                reason_texts=("관심 분야와 일치합니다.", "방문 목적과 관련 있습니다."),
                evidence_refs=("profile:category", "catalog:approved"),
                eligibility_evaluation_id="eligibility-001",
                score_policy_version="consumer-score-v1.0",
                score_fingerprint="score-fingerprint-001",
                recommendation_session_id=str(SESSION_ID),
            ),
        ),
        statuses=(
            BatchProfileStatus(
                source_record_id="VISITOR-001",
                status="SUCCESS",
                result_count=1,
                error_code=None,
                recommendation_session_id=str(SESSION_ID),
                policy_version="consumer-score-v1.0",
                ranking_version="ranking-v1",
                explanation_version="explanation-v1",
                candidate_count=4,
                filtered_count=2,
            ),
            BatchProfileStatus(
                source_record_id="VISITOR-002",
                status="SKIPPED",
                result_count=0,
                error_code="PERSONALIZATION_DISABLED",
                recommendation_session_id=None,
                policy_version=None,
                ranking_version=None,
                explanation_version=None,
                candidate_count=0,
                filtered_count=0,
            ),
        ),
    )


def test_workbook_contains_formula_summary_safe_matches_and_statuses() -> None:
    payload = render_batch_matching_workbook(_report())
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        assert workbook.sheetnames == ["요약", "매칭결과", "처리상태"]
        assert workbook["요약"]["B5"].value == "=COUNTA('처리상태'!A2:A101)"
        assert (
            workbook["요약"]["B6"].value == "=COUNTIF('처리상태'!B2:B101,\"SUCCESS\")"
        )
        assert workbook["매칭결과"]["A2"].value == "VISITOR-001"
        assert workbook["매칭결과"]["C2"].value == "샘플 양조장"
        assert workbook["매칭결과"]["F2"].value == 87.25
        assert workbook["처리상태"]["D3"].value == "PERSONALIZATION_DISABLED"
        all_text = " ".join(
            str(cell.value)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        assert "phone" not in all_text.casefold()
        assert "email" not in all_text.casefold()
    finally:
        workbook.close()


@pytest.mark.asyncio
async def test_batch_matching_uses_canonical_orchestrator_and_isolates_failures(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def fake_resolve(*args, **kwargs):
        del args, kwargs
        return {
            "VISITOR-001": SourceProfileRef("VISITOR-001", PROFILE_ID, USER_ID),
            "VISITOR-002": SourceProfileRef("VISITOR-002", uuid.uuid4(), uuid.uuid4()),
        }

    monkeypatch.setattr(export_service, "_resolve_source_profiles", fake_resolve)

    candidate = MatchCandidate(
        object_type="EXHIBITOR",
        object_id=EXHIBITOR_ID,
        recommendable_id=uuid.uuid4(),
        exhibitor_id=EXHIBITOR_ID,
        participation_id=uuid.uuid4(),
        public_object_id=str(EXHIBITOR_ID),
        payload={"company_name": "샘플 양조장"},
    )
    candidate.rank = 1
    candidate.final_score = 0.8725
    candidate.filter_evaluation_id = "eligibility-001"
    candidate.directional_policy_version = "consumer-score-v1.0"
    candidate.directional_score_fingerprint = "score-fingerprint-001"
    candidate.reasons = [
        MatchReasonDraft(
            code="CATEGORY_MATCH",
            text="관심 분야와 일치합니다.",
            evidence_refs=["catalog:approved"],
            contribution_score=0.4,
            generated_by="TEMPLATE",
        )
    ]

    class FakeOrchestrator:
        async def generate(self, db, *, subject, **kwargs):
            del db, kwargs
            if subject.profile_id != PROFILE_ID:
                raise personalization_disabled()
            return RecommendationOutcome(
                recommendation_session_id=SESSION_ID,
                generated_at=datetime.now(UTC),
                expires_at=datetime.now(UTC),
                profile_version=1,
                ranking_version="ranking-v1",
                explanation_version="explanation-v1",
                policy_version="consumer-score-v1.0",
                items=[candidate],
                trace=PipelineTrace(
                    candidate_count=4,
                    filtered_count=2,
                    result_count=1,
                ),
            )

    class FakeDb:
        def __init__(self) -> None:
            self.rollback_count = 0

        async def rollback(self) -> None:
            self.rollback_count += 1

    fake_db = FakeDb()
    report = await run_batch_matching(
        fake_db,  # type: ignore[arg-type]
        tenant_id=TENANT_ID,
        event_id=EVENT_ID,
        source_system_code="EXCEL_UPLOAD",
        source_record_ids=("VISITOR-001", "MISSING", "VISITOR-002"),
        top_n=5,
        orchestrator=FakeOrchestrator(),  # type: ignore[arg-type]
    )

    assert [status.status for status in report.statuses] == [
        "SUCCESS",
        "SKIPPED",
        "SKIPPED",
    ]
    assert report.statuses[1].error_code == "EXCEL_MATCH_PROFILE_NOT_FOUND"
    assert report.statuses[2].error_code == "PERSONALIZATION_DISABLED"
    assert len(report.matches) == 1
    assert report.matches[0].score_policy_version == "consumer-score-v1.0"
    assert fake_db.rollback_count == 1


def test_request_rejects_duplicate_source_record_ids() -> None:
    with pytest.raises(ValidationError):
        ExcelMatchExportRequest(
            tenant_id=TENANT_ID,
            event_id=EVENT_ID,
            source_record_ids=["VISITOR-001", "VISITOR-001"],
        )


def test_endpoint_returns_xlsx_with_schema_header(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    async def fake_db():
        yield object()

    async def fake_run(*args, **kwargs):
        del args, kwargs
        return _report()

    monkeypatch.setattr(imports_router, "run_batch_matching", fake_run)
    app.dependency_overrides[get_db] = fake_db
    try:
        response = TestClient(app).post(
            "/api/v1/admin/imports/excel/matches",
            json={
                "tenant_id": str(TENANT_ID),
                "event_id": str(EVENT_ID),
                "source_system_code": "EXCEL_UPLOAD",
                "source_record_ids": ["VISITOR-001", "VISITOR-002"],
                "top_n": 5,
            },
        )
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert response.headers["x-workbook-schema"] == EXCEL_BATCH_MATCH_SCHEMA_VERSION
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["요약", "매칭결과", "처리상태"]
    finally:
        workbook.close()


def test_openapi_publishes_xlsx_response_contract() -> None:
    operation = app.openapi()["paths"]["/api/v1/admin/imports/excel/matches"]["post"]
    assert XLSX_MEDIA_TYPE in operation["responses"]["200"]["content"]
    assert operation["requestBody"]["content"]["application/json"]["schema"] == {
        "$ref": "#/components/schemas/ExcelMatchExportRequest"
    }
