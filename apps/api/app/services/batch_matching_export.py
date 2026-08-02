"""Operator-only batch recommendation export for imported preregistrants.

This adapter never scores workbook rows directly. It resolves opaque external IDs to canonical
profiles, executes the same consent checks, approved-catalog retrieval, Hard Filter, matching
facade, explanation policy, and result persistence as the registered web, then projects the safe
result subset to XLSX.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.integration import ExternalReference, SourceSystem
from app.models.profile import UserProfile
from app.services.matching.errors import RecommendationError
from app.services.matching.orchestrator import (
    RecommendationOrchestrator,
    recommendation_orchestrator,
)
from app.services.matching.types import SubjectContext

EXCEL_BATCH_MATCH_SCHEMA_VERSION = "meet-ai-batch-match-export-v1.0"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class BatchMatchingExportError(ValueError):
    """Safe request-level export error."""

    def __init__(self, code: str, http_status: int):
        self.code = code
        self.http_status = http_status
        super().__init__(code)


@dataclass(frozen=True, slots=True)
class SourceProfileRef:
    source_record_id: str
    profile_id: uuid.UUID
    user_id: uuid.UUID | None


@dataclass(frozen=True, slots=True)
class BatchMatchEntry:
    source_record_id: str
    rank: int
    exhibitor_name: str
    exhibitor_id: str
    object_id: str
    score_100: float
    match_level: str
    recommended_action: str
    reason_codes: tuple[str, ...]
    reason_texts: tuple[str, ...]
    evidence_refs: tuple[str, ...]
    eligibility_evaluation_id: str
    score_policy_version: str
    score_fingerprint: str
    recommendation_session_id: str


@dataclass(frozen=True, slots=True)
class BatchProfileStatus:
    source_record_id: str
    status: str
    result_count: int
    error_code: str | None
    recommendation_session_id: str | None
    policy_version: str | None
    ranking_version: str | None
    explanation_version: str | None
    candidate_count: int
    filtered_count: int


@dataclass(frozen=True, slots=True)
class BatchMatchingReport:
    schema_version: str
    tenant_id: uuid.UUID
    event_id: uuid.UUID
    source_system_code: str
    top_n: int
    generated_at: datetime
    matches: tuple[BatchMatchEntry, ...]
    statuses: tuple[BatchProfileStatus, ...]


async def _resolve_source_profiles(
    db: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    event_id: uuid.UUID,
    source_system_code: str,
    source_record_ids: Sequence[str],
) -> Mapping[str, SourceProfileRef]:
    source_system_id = (
        await db.execute(
            select(SourceSystem.source_system_id).where(
                SourceSystem.tenant_id == tenant_id,
                SourceSystem.system_code == source_system_code,
                SourceSystem.active.is_(True),
            )
        )
    ).scalar_one_or_none()
    if source_system_id is None:
        raise BatchMatchingExportError(
            "EXCEL_MATCH_SOURCE_SYSTEM_NOT_FOUND", http_status=404
        )

    rows = (
        await db.execute(
            select(
                ExternalReference.external_id,
                UserProfile.profile_id,
                UserProfile.user_id,
            )
            .join(
                UserProfile,
                UserProfile.profile_id == ExternalReference.internal_id,
            )
            .where(
                ExternalReference.source_system_id == source_system_id,
                ExternalReference.object_type == "VISITOR",
                ExternalReference.external_id.in_(tuple(source_record_ids)),
                ExternalReference.sync_status == "SYNCED",
                UserProfile.tenant_id == tenant_id,
                UserProfile.event_id == event_id,
                UserProfile.deleted_at.is_(None),
            )
        )
    ).all()
    return {
        row.external_id: SourceProfileRef(
            source_record_id=row.external_id,
            profile_id=row.profile_id,
            user_id=row.user_id,
        )
        for row in rows
    }


def _safe_tuple(values: Sequence[str | None]) -> tuple[str, ...]:
    return tuple(sorted({value.strip() for value in values if value and value.strip()}))


async def run_batch_matching(
    db: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    event_id: uuid.UUID,
    source_system_code: str,
    source_record_ids: Sequence[str],
    top_n: int,
    orchestrator: RecommendationOrchestrator = recommendation_orchestrator,
) -> BatchMatchingReport:
    """Run one canonical EXHIBITOR recommendation for each imported source ID."""

    generated_at = datetime.now(UTC)
    resolved = await _resolve_source_profiles(
        db,
        tenant_id=tenant_id,
        event_id=event_id,
        source_system_code=source_system_code,
        source_record_ids=source_record_ids,
    )
    matches: list[BatchMatchEntry] = []
    statuses: list[BatchProfileStatus] = []

    for source_record_id in source_record_ids:
        profile = resolved.get(source_record_id)
        if profile is None:
            statuses.append(
                BatchProfileStatus(
                    source_record_id=source_record_id,
                    status="SKIPPED",
                    result_count=0,
                    error_code="EXCEL_MATCH_PROFILE_NOT_FOUND",
                    recommendation_session_id=None,
                    policy_version=None,
                    ranking_version=None,
                    explanation_version=None,
                    candidate_count=0,
                    filtered_count=0,
                )
            )
            continue

        subject = SubjectContext(
            tenant_id=tenant_id,
            event_id=event_id,
            profile_id=profile.profile_id,
            visit_session_id=None,
            user_id=profile.user_id,
            guest_session_id=None,
            request_id=f"batch-match-{uuid.uuid4()}",
            idempotency_key=None,
            server_time=generated_at,
        )
        try:
            outcome = await orchestrator.generate(
                db,
                subject=subject,
                recommendation_type="EXHIBITOR",
                context_input={},
                limit=top_n,
            )
        except RecommendationError as exc:
            # A failed pipeline may already have staged audit/filter rows before raising
            # NO_CANDIDATE. Roll them back so a later successful profile cannot commit another
            # profile's partial work as a side effect of this batch.
            await db.rollback()
            statuses.append(
                BatchProfileStatus(
                    source_record_id=source_record_id,
                    status="SKIPPED",
                    result_count=0,
                    error_code=exc.code,
                    recommendation_session_id=None,
                    policy_version=None,
                    ranking_version=None,
                    explanation_version=None,
                    candidate_count=0,
                    filtered_count=0,
                )
            )
            continue

        session_id = str(outcome.recommendation_session_id)
        for candidate in outcome.items:
            reason_codes = _safe_tuple([reason.code for reason in candidate.reasons])
            reason_texts = tuple(
                reason.text.strip()
                for reason in candidate.reasons
                if reason.text.strip()
            )
            evidence_refs = _safe_tuple(
                [ref for reason in candidate.reasons for ref in reason.evidence_refs]
            )
            matches.append(
                BatchMatchEntry(
                    source_record_id=source_record_id,
                    rank=candidate.rank,
                    exhibitor_name=str(
                        candidate.payload.get("company_name")
                        or candidate.payload.get("exhibitor_name")
                        or ""
                    ),
                    exhibitor_id=str(candidate.exhibitor_id or ""),
                    object_id=candidate.public_object_id,
                    score_100=round(float(candidate.final_score) * 100, 4),
                    match_level=candidate.match_level(),
                    recommended_action=candidate.recommended_action,
                    reason_codes=reason_codes,
                    reason_texts=reason_texts,
                    evidence_refs=evidence_refs,
                    eligibility_evaluation_id=str(candidate.filter_evaluation_id or ""),
                    score_policy_version=str(
                        candidate.directional_policy_version or outcome.policy_version
                    ),
                    score_fingerprint=str(
                        candidate.directional_score_fingerprint or ""
                    ),
                    recommendation_session_id=session_id,
                )
            )
        statuses.append(
            BatchProfileStatus(
                source_record_id=source_record_id,
                status="SUCCESS",
                result_count=len(outcome.items),
                error_code=None,
                recommendation_session_id=session_id,
                policy_version=outcome.policy_version,
                ranking_version=outcome.ranking_version,
                explanation_version=outcome.explanation_version,
                candidate_count=outcome.trace.candidate_count,
                filtered_count=outcome.trace.filtered_count,
            )
        )

    return BatchMatchingReport(
        schema_version=EXCEL_BATCH_MATCH_SCHEMA_VERSION,
        tenant_id=tenant_id,
        event_id=event_id,
        source_system_code=source_system_code,
        top_n=top_n,
        generated_at=generated_at,
        matches=tuple(matches),
        statuses=tuple(statuses),
    )


_NAVY = "17324D"
_TEAL = "0F766E"
_PALE_TEAL = "DDF4EF"
_PALE_AMBER = "FFF4CC"
_PALE_RED = "FDE8E7"
_WHITE = "FFFFFF"
_TEXT = "23313F"
_MUTED = "607080"
_BORDER = Side(style="thin", color="D8E0E8")


def _style_header(row: Sequence[Any]) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.font = Font(color=_WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=_BORDER)


def _configure_data_sheet(
    sheet: Any,
    *,
    widths: Sequence[int],
    table_name: str,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 28
    _style_header(sheet[1])
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    if sheet.max_row > 1:
        table = Table(
            displayName=table_name,
            ref=f"A1:{sheet.cell(1, sheet.max_column).column_letter}{sheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    sheet.auto_filter.ref = (
        f"A1:{sheet.cell(1, sheet.max_column).column_letter}{max(sheet.max_row, 1)}"
    )
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0


def render_batch_matching_workbook(report: BatchMatchingReport) -> bytes:
    """Render a safe operator workbook; no direct identity or raw profile values are included."""

    workbook = Workbook()
    workbook.properties.title = "Meet AI 사전등록자별 참여기업 매칭 결과"
    workbook.properties.subject = report.schema_version
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    summary = workbook.active
    summary.title = "요약"
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F2")
    summary["A1"] = "사전등록자별 Top-N 참여기업 매칭 결과"
    summary["A1"].fill = PatternFill("solid", fgColor=_NAVY)
    summary["A1"].font = Font(color=_WHITE, bold=True, size=18)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 28
    summary.row_dimensions[2].height = 12

    summary_rows = [
        ("스키마 버전", report.schema_version),
        ("생성 시각(UTC)", report.generated_at.replace(tzinfo=None)),
        ("요청 사전등록자", "=COUNTA('처리상태'!A2:A101)"),
        ("성공", "=COUNTIF('처리상태'!B2:B101,\"SUCCESS\")"),
        ("확인 필요", "=B5-B6"),
        ("생성 매칭 수", "=COUNTA('매칭결과'!A2:A1001)"),
        ("Top-N", report.top_n),
        ("원천 시스템", report.source_system_code),
        ("행사 ID", str(report.event_id)),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=3):
        summary.cell(row_index, 1, label)
        summary.cell(row_index, 2, value)
        summary.cell(row_index, 1).fill = PatternFill("solid", fgColor="E8EEF4")
        summary.cell(row_index, 1).font = Font(bold=True, color=_TEXT)
        summary.cell(row_index, 2).font = Font(color=_TEXT)
        summary.cell(row_index, 1).border = Border(bottom=_BORDER)
        summary.cell(row_index, 2).border = Border(bottom=_BORDER)
    summary["B4"].number_format = "yyyy-mm-dd hh:mm:ss"
    summary["B5"].number_format = "#,##0"
    summary["B6"].number_format = "#,##0"
    summary["B7"].number_format = "#,##0"
    summary["B8"].number_format = "#,##0"
    summary["B9"].number_format = "#,##0"
    summary.merge_cells("A14:F16")
    summary["A14"] = (
        "이 파일은 운영자 검토용입니다. 기존 개인화 동의·연령확인·승인 카탈로그·Hard Filter를 "
        "통과한 결과만 매칭결과 시트에 포함합니다. 확인 필요 건은 조건을 임의 완화하지 않습니다."
    )
    summary["A14"].fill = PatternFill("solid", fgColor=_PALE_AMBER)
    summary["A14"].font = Font(color=_TEXT)
    summary["A14"].alignment = Alignment(wrap_text=True, vertical="top")
    summary.column_dimensions["A"].width = 25
    summary.column_dimensions["B"].width = 42
    for column in ("C", "D", "E", "F"):
        summary.column_dimensions[column].width = 14

    results = workbook.create_sheet("매칭결과")
    result_headers = (
        "source_record_id",
        "rank",
        "exhibitor_name",
        "exhibitor_id",
        "object_id",
        "score_100",
        "match_level",
        "recommended_action",
        "reason_codes",
        "reason_texts",
        "evidence_refs",
        "eligibility_evaluation_id",
        "score_policy_version",
        "score_fingerprint",
        "recommendation_session_id",
    )
    results.append(result_headers)
    for entry in report.matches:
        results.append(
            (
                entry.source_record_id,
                entry.rank,
                entry.exhibitor_name,
                entry.exhibitor_id,
                entry.object_id,
                entry.score_100,
                entry.match_level,
                entry.recommended_action,
                ";".join(entry.reason_codes),
                " | ".join(entry.reason_texts),
                ";".join(entry.evidence_refs),
                entry.eligibility_evaluation_id,
                entry.score_policy_version,
                entry.score_fingerprint,
                entry.recommendation_session_id,
            )
        )
    _configure_data_sheet(
        results,
        widths=(20, 8, 26, 38, 38, 12, 14, 22, 28, 60, 38, 38, 24, 64, 38),
        table_name="BatchMatchResults",
    )
    results.column_dimensions["B"].width = 8
    for cell in results["F"][1:]:
        cell.number_format = "0.00"
    for row in results.iter_rows(min_row=2):
        row[9].alignment = Alignment(wrap_text=True, vertical="top")

    statuses = workbook.create_sheet("처리상태")
    status_headers = (
        "source_record_id",
        "status",
        "result_count",
        "error_code",
        "recommendation_session_id",
        "policy_version",
        "ranking_version",
        "explanation_version",
        "candidate_count",
        "filtered_count",
    )
    statuses.append(status_headers)
    for status in report.statuses:
        statuses.append(
            (
                status.source_record_id,
                status.status,
                status.result_count,
                status.error_code,
                status.recommendation_session_id,
                status.policy_version,
                status.ranking_version,
                status.explanation_version,
                status.candidate_count,
                status.filtered_count,
            )
        )
    _configure_data_sheet(
        statuses,
        widths=(20, 14, 14, 34, 38, 24, 24, 24, 16, 16),
        table_name="BatchMatchStatuses",
    )
    for row_index in range(2, statuses.max_row + 1):
        fill = (
            _PALE_TEAL if statuses.cell(row_index, 2).value == "SUCCESS" else _PALE_RED
        )
        statuses.cell(row_index, 2).fill = PatternFill("solid", fgColor=fill)
        statuses.cell(row_index, 2).font = Font(bold=True, color=_TEXT)

    return_buffer = io.BytesIO()
    workbook.save(return_buffer)
    workbook.close()
    return return_buffer.getvalue()
