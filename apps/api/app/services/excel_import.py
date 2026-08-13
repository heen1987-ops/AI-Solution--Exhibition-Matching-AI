"""Fail-closed XLSX adapter for the existing canonical JSON import contracts.

The workbook is an input transport, never a second source of matching truth. Human-readable
ontology codes are resolved against the published catalog and converted to the same immutable UUID
references used by the JSON import endpoints. No workbook value is logged or included in file-level
errors.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Literal

from meet_ai.ontology.catalog import Catalog, load_catalog, stable_uuid
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from pydantic import ValidationError

from app.schemas.imports import (
    ExhibitorImportRow,
    ImportRowError,
    TaxonomyAttributeRef,
    VisitorImportRow,
)

EXCEL_IMPORT_SCHEMA_VERSION = "meet-ai-excel-import-v1.0"
MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 1_000
MAX_ROWS_PER_SHEET = 1_000

VISITOR_SHEET = "사전등록자"
EXHIBITOR_SHEET = "참여기업"

VISITOR_HEADERS = (
    "source_record_id",
    "source_updated_at",
    "name",
    "phone",
    "email",
    "home_region",
    "age_band",
    "interest_codes",
    "visit_goal_codes",
    "attribution_channels",
    "consent_registration",
    "consent_marketing",
    "age_19_plus",
)
EXHIBITOR_HEADERS = (
    "source_record_id",
    "source_updated_at",
    "legal_name",
    "display_name",
    "website_url",
    "story",
    "manager_name",
    "manager_phone",
    "manager_email",
    "region_code",
    "category_codes",
    "promotion_description",
)

_VISITOR_INTEREST_TYPES = frozenset(
    {
        "ALCOHOL_LEVEL",
        "AROMA",
        "BOOTH_SERVICE",
        "BUSINESS_GOAL",
        "CHANNEL",
        "INGREDIENT",
        "PRICE_BAND",
        "PRODUCT_CATEGORY",
        "PRODUCT_FEATURE",
        "REGION",
        "SUPPLY_CAPACITY",
        "TASTE",
        "TRADE_TYPE",
        "USE_CASE",
        "VISIT_GOAL",
    }
)
_VISITOR_GOAL_TYPES = frozenset({"VISIT_GOAL", "BUSINESS_GOAL"})
_EXHIBITOR_CATEGORY_TYPES = frozenset({"PRODUCT_CATEGORY", "TRADE_TYPE"})


class ExcelImportValidationError(ValueError):
    """Safe file-level error suitable for returning to an operator."""

    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


class _RowValidationError(ValueError):
    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


@dataclass(frozen=True, slots=True)
class ParsedSheet:
    sheet_name: Literal["사전등록자", "참여기업"]
    total_rows: int
    rows: tuple[VisitorImportRow | ExhibitorImportRow, ...]
    row_numbers: tuple[int, ...]
    errors: tuple[ImportRowError, ...]


@dataclass(frozen=True, slots=True)
class ParsedExcelImport:
    schema_version: str
    visitors: ParsedSheet
    exhibitors: ParsedSheet


def _validate_archive(data: bytes) -> None:
    if not data or len(data) > MAX_XLSX_BYTES:
        raise ExcelImportValidationError(
            "EXCEL_FILE_SIZE_INVALID", "XLSX 파일은 1바이트 이상 5MiB 이하여야 합니다."
        )
    if not data.startswith(b"PK"):
        raise ExcelImportValidationError(
            "EXCEL_FILE_INVALID", "유효한 .xlsx 파일이 아닙니다."
        )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise ExcelImportValidationError(
                    "EXCEL_ARCHIVE_LIMIT_EXCEEDED",
                    "XLSX 내부 파일 수가 허용 한도를 초과했습니다.",
                )
            if any(
                ".." in info.filename.replace("\\", "/").split("/") for info in infos
            ):
                raise ExcelImportValidationError(
                    "EXCEL_FILE_INVALID", "XLSX 내부 경로가 유효하지 않습니다."
                )
            if any(info.filename.lower().endswith("vbaproject.bin") for info in infos):
                raise ExcelImportValidationError(
                    "EXCEL_MACRO_NOT_ALLOWED",
                    "매크로가 포함된 파일은 허용되지 않습니다.",
                )
            if sum(info.file_size for info in infos) > MAX_UNCOMPRESSED_BYTES:
                raise ExcelImportValidationError(
                    "EXCEL_ARCHIVE_LIMIT_EXCEEDED",
                    "XLSX 압축 해제 크기가 허용 한도를 초과했습니다.",
                )
    except zipfile.BadZipFile as exc:
        raise ExcelImportValidationError(
            "EXCEL_FILE_INVALID", "유효한 .xlsx 파일이 아닙니다."
        ) from exc


def _text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _required_text(value: Any, field_name: str) -> str:
    normalized = _text(value)
    if normalized is None:
        raise _RowValidationError(
            "EXCEL_REQUIRED_VALUE_MISSING", f"필수 열 {field_name} 값이 없습니다."
        )
    return normalized


def _timestamp(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        raw = _required_text(value, "source_updated_at")
        try:
            parsed = datetime.fromisoformat(raw)
        except ValueError as exc:
            raise _RowValidationError(
                "EXCEL_DATETIME_INVALID",
                "source_updated_at은 ISO-8601 날짜시간이어야 합니다.",
            ) from exc
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=UTC)
    return parsed


def _boolean(value: Any, field_name: str, *, default: bool | None = None) -> bool:
    if value is None or (isinstance(value, str) and not value.strip()):
        if default is not None:
            return default
        raise _RowValidationError(
            "EXCEL_REQUIRED_VALUE_MISSING", f"필수 열 {field_name} 값이 없습니다."
        )
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().casefold()
    if normalized in {"true", "y", "yes", "1", "예", "동의"}:
        return True
    if normalized in {"false", "n", "no", "0", "아니오", "미동의"}:
        return False
    raise _RowValidationError(
        "EXCEL_BOOLEAN_INVALID",
        f"{field_name}은 TRUE/FALSE 또는 Y/N 값이어야 합니다.",
    )


def _string_list(value: Any) -> list[str]:
    raw = _text(value)
    if raw is None:
        return []
    return list(dict.fromkeys(part.strip() for part in raw.split(";") if part.strip()))


def _taxonomy_refs(
    value: Any,
    *,
    catalog: Catalog,
    field_name: str,
    allowed_types: frozenset[str],
) -> list[TaxonomyAttributeRef]:
    refs: list[TaxonomyAttributeRef] = []
    taxonomy_version_id = stable_uuid("taxonomy-version", catalog.version)
    for code in _string_list(value):
        item = catalog.by_code.get(code)
        if item is None:
            raise _RowValidationError(
                "EXCEL_ONTOLOGY_CODE_UNKNOWN",
                f"{field_name}에 게시되지 않은 온톨로지 코드가 있습니다.",
            )
        if item.get("assignable", True) is False:
            raise _RowValidationError(
                "EXCEL_ONTOLOGY_CODE_NOT_ASSIGNABLE",
                f"{field_name}에 선택할 수 없는 상위 코드가 있습니다.",
            )
        if item.get("concept_type") not in allowed_types:
            raise _RowValidationError(
                "EXCEL_ONTOLOGY_TYPE_INVALID",
                f"{field_name}에 허용되지 않은 유형의 코드가 있습니다.",
            )
        refs.append(
            TaxonomyAttributeRef(
                taxonomy_version_id=taxonomy_version_id,
                concept_id=stable_uuid("concept", code),
                attribute_code=code,
                value=True,
            )
        )
    return refs


def _source_id(row: dict[str, Any]) -> str | None:
    value = _text(row.get("source_record_id"))
    return value[:255] if value else None


def _visitor_row(row: dict[str, Any], catalog: Catalog) -> VisitorImportRow:
    return VisitorImportRow(
        source_record_id=_required_text(
            row.get("source_record_id"), "source_record_id"
        ),
        source_updated_at=_timestamp(row.get("source_updated_at")),
        name=_text(row.get("name")),
        phone=_text(row.get("phone")),
        email=_text(row.get("email")),
        home_region=_text(row.get("home_region")),
        age_band=_text(row.get("age_band")),
        interest_categories=_taxonomy_refs(
            row.get("interest_codes"),
            catalog=catalog,
            field_name="interest_codes",
            allowed_types=_VISITOR_INTEREST_TYPES,
        ),
        visit_goals=_taxonomy_refs(
            row.get("visit_goal_codes"),
            catalog=catalog,
            field_name="visit_goal_codes",
            allowed_types=_VISITOR_GOAL_TYPES,
        ),
        attribution_channels=_string_list(row.get("attribution_channels")),
        consent_registration=_boolean(
            row.get("consent_registration"), "consent_registration"
        ),
        consent_marketing=_boolean(
            row.get("consent_marketing"), "consent_marketing", default=False
        ),
        age_19_plus=_boolean(row.get("age_19_plus"), "age_19_plus"),
    )


def _exhibitor_row(row: dict[str, Any], catalog: Catalog) -> ExhibitorImportRow:
    region_refs = _taxonomy_refs(
        row.get("region_code"),
        catalog=catalog,
        field_name="region_code",
        allowed_types=frozenset({"REGION"}),
    )
    if len(region_refs) > 1:
        raise _RowValidationError(
            "EXCEL_SINGLE_VALUE_REQUIRED",
            "region_code에는 코드 하나만 입력해야 합니다.",
        )
    return ExhibitorImportRow(
        source_record_id=_required_text(
            row.get("source_record_id"), "source_record_id"
        ),
        source_updated_at=_timestamp(row.get("source_updated_at")),
        legal_name=_required_text(row.get("legal_name"), "legal_name"),
        display_name=_text(row.get("display_name")),
        website_url=_text(row.get("website_url")),
        story=_text(row.get("story")),
        manager_name=_text(row.get("manager_name")),
        manager_phone=_text(row.get("manager_phone")),
        manager_email=_text(row.get("manager_email")),
        region=region_refs[0] if region_refs else None,
        categories=_taxonomy_refs(
            row.get("category_codes"),
            catalog=catalog,
            field_name="category_codes",
            allowed_types=_EXHIBITOR_CATEGORY_TYPES,
        ),
        promotion_description=_text(row.get("promotion_description")),
    )


def _safe_pydantic_error(exc: ValidationError) -> tuple[str, str]:
    first = exc.errors(include_url=False, include_context=False, include_input=False)[0]
    field = ".".join(str(part) for part in first.get("loc", ())) or "row"
    return (
        "EXCEL_ROW_SCHEMA_INVALID",
        f"{field} 열의 형식 또는 길이가 유효하지 않습니다.",
    )


def _parse_sheet(
    worksheet,
    *,
    sheet_name: Literal["사전등록자", "참여기업"],
    headers: tuple[str, ...],
    catalog: Catalog,
) -> ParsedSheet:
    # Some valid XLSX producers omit the worksheet dimension metadata. In
    # read-only mode openpyxl then leaves max_row/max_column unset until the
    # dimension is calculated from the worksheet XML.
    if worksheet.max_row is None or worksheet.max_column is None:
        worksheet.calculate_dimension(force=True)
    if worksheet.max_row > 100_000 or worksheet.max_column > 100:
        raise ExcelImportValidationError(
            "EXCEL_SHEET_DIMENSION_INVALID",
            f"{sheet_name} 시트 범위가 비정상적으로 큽니다.",
        )
    if worksheet.max_column > len(headers):
        raise ExcelImportValidationError(
            "EXCEL_HEADERS_INVALID",
            f"{sheet_name} 시트에 템플릿에 없는 추가 열이 있습니다.",
        )
    header_cells: tuple[Cell, ...] = next(
        worksheet.iter_rows(min_row=1, max_row=1, max_col=len(headers))
    )
    if any(cell.data_type == "f" for cell in header_cells):
        raise ExcelImportValidationError(
            "EXCEL_FORMULA_NOT_ALLOWED",
            f"{sheet_name} 시트에는 수식을 사용할 수 없습니다.",
        )
    observed_headers = tuple(_text(cell.value) or "" for cell in header_cells)
    if observed_headers != headers:
        raise ExcelImportValidationError(
            "EXCEL_HEADERS_INVALID",
            f"{sheet_name} 시트의 열 이름 또는 순서가 템플릿과 다릅니다.",
        )

    parsed: list[VisitorImportRow | ExhibitorImportRow] = []
    parsed_row_numbers: list[int] = []
    errors: list[ImportRowError] = []
    seen_source_ids: set[str] = set()
    total_rows = 0
    parser = _visitor_row if sheet_name == VISITOR_SHEET else _exhibitor_row

    for excel_row_index, cells in enumerate(
        worksheet.iter_rows(min_row=2, max_col=len(headers)), start=2
    ):
        if all(_text(cell.value) is None for cell in cells):
            continue
        total_rows += 1
        if total_rows > MAX_ROWS_PER_SHEET:
            raise ExcelImportValidationError(
                "EXCEL_ROW_LIMIT_EXCEEDED",
                f"{sheet_name} 시트는 최대 1,000행까지 처리할 수 있습니다.",
            )
        values = dict(zip(headers, (cell.value for cell in cells), strict=True))
        source_record_id = _source_id(values)
        if any(cell.data_type == "f" for cell in cells):
            errors.append(
                ImportRowError(
                    row_index=excel_row_index,
                    source_record_id=None,
                    error_code="EXCEL_FORMULA_NOT_ALLOWED",
                    message="입력 행에는 수식을 사용할 수 없습니다.",
                )
            )
            continue
        if source_record_id and source_record_id in seen_source_ids:
            errors.append(
                ImportRowError(
                    row_index=excel_row_index,
                    source_record_id=None,
                    error_code="EXCEL_SOURCE_ID_DUPLICATE",
                    message="같은 시트에 중복된 source_record_id가 있습니다.",
                )
            )
            continue
        try:
            normalized = parser(values, catalog)
        except _RowValidationError as exc:
            errors.append(
                ImportRowError(
                    row_index=excel_row_index,
                    source_record_id=None,
                    error_code=exc.code,
                    message=exc.message,
                )
            )
        except ValidationError as exc:
            error_code, message = _safe_pydantic_error(exc)
            errors.append(
                ImportRowError(
                    row_index=excel_row_index,
                    source_record_id=None,
                    error_code=error_code,
                    message=message,
                )
            )
        else:
            parsed.append(normalized)
            parsed_row_numbers.append(excel_row_index)
            seen_source_ids.add(normalized.source_record_id)

    return ParsedSheet(
        sheet_name,
        total_rows,
        tuple(parsed),
        tuple(parsed_row_numbers),
        tuple(errors),
    )


def parse_excel_import(
    data: bytes, *, catalog: Catalog | None = None
) -> ParsedExcelImport:
    """Parse an XLSX file into canonical visitor/exhibitor import rows without persistence."""

    _validate_archive(data)
    resolved_catalog = catalog or load_catalog()
    try:
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
    except Exception as exc:
        raise ExcelImportValidationError(
            "EXCEL_FILE_INVALID", "XLSX 파일을 읽을 수 없습니다."
        ) from exc
    try:
        if any(
            workbook[name].sheet_state != "visible"
            for name in (VISITOR_SHEET, EXHIBITOR_SHEET)
            if name in workbook.sheetnames
        ):
            raise ExcelImportValidationError(
                "EXCEL_INPUT_SHEET_HIDDEN", "입력 시트를 숨길 수 없습니다."
            )
        visitors = (
            _parse_sheet(
                workbook[VISITOR_SHEET],
                sheet_name=VISITOR_SHEET,
                headers=VISITOR_HEADERS,
                catalog=resolved_catalog,
            )
            if VISITOR_SHEET in workbook.sheetnames
            else ParsedSheet(VISITOR_SHEET, 0, (), (), ())
        )
        exhibitors = (
            _parse_sheet(
                workbook[EXHIBITOR_SHEET],
                sheet_name=EXHIBITOR_SHEET,
                headers=EXHIBITOR_HEADERS,
                catalog=resolved_catalog,
            )
            if EXHIBITOR_SHEET in workbook.sheetnames
            else ParsedSheet(EXHIBITOR_SHEET, 0, (), (), ())
        )
        if visitors.total_rows + exhibitors.total_rows == 0:
            raise ExcelImportValidationError(
                "EXCEL_INPUT_EMPTY",
                "사전등록자 또는 참여기업 입력 행이 한 건 이상 필요합니다.",
            )
        return ParsedExcelImport(
            EXCEL_IMPORT_SCHEMA_VERSION,
            visitors,
            exhibitors,
        )
    finally:
        workbook.close()
